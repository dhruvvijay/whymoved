"""Load a messy file, show the guessed roles and validation report, then explain a change.

Usage:
    python scripts/demo_ingest.py                  # builds a messy SYNTHETIC .xlsx, loads it
    python scripts/demo_ingest.py my_file.xlsx          # your own CSV or Excel file
    python scripts/demo_ingest.py my_file.csv --metric revenue
    python scripts/demo_ingest.py my_file.csv --a 2026-01-05:2026-02-01 --b 2026-02-02:2026-03-01

Only additive metrics (revenue, orders, tickets) are explained here; ratios arrive in M5.
"""

import argparse
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from whymoved.decompose import additive  # noqa: E402
from whymoved.generator import make_scenario  # noqa: E402
from whymoved.ingest import DataError, default_periods, prepare  # noqa: E402
from whymoved.trust import explain_with_trust  # noqa: E402


def build_messy_sample() -> tuple[Path, str]:  # also used by report.py
    """A SYNTHETIC report-style workbook: title lines, text money, case variants, totals, a gap."""
    sc = make_scenario("single", 7)
    d = sc.data.copy()
    d = d[~d["date"].isin(pd.to_datetime(["2026-02-10", "2026-02-11"]))]  # two missing days
    d = (
        d.groupby(["date", "region", "device"], as_index=False)[
            ["orders", "revenue", "budget_revenue"]
        ].sum()
        .reset_index(drop=True)
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Quarterly sales report (SYNTHETIC)"
    ws["A2"] = "Prepared by finance"
    headers = ["Date", "Region", "Device", "Orders", "Revenue", "Budget_Revenue"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=4, column=j, value=h)
    row = 5
    for i, rec in enumerate(d.itertuples(index=False)):
        region = rec.region.upper() if i % 37 == 0 else rec.region
        ws.cell(row=row, column=1, value=rec.date.strftime("%d/%m/%Y"))
        ws.cell(row=row, column=2, value=region)
        ws.cell(row=row, column=3, value=rec.device)
        ws.cell(row=row, column=4, value=int(rec.orders))
        ws.cell(row=row, column=5, value=f"₹{rec.revenue:,.2f}")
        ws.cell(row=row, column=6, value=f"₹{rec.budget_revenue:,.2f}")
        row += 1
    ws.cell(row=row, column=2, value="Grand Total")
    ws.cell(row=row, column=5, value=f"₹{d['revenue'].sum():,.2f}")
    path = Path(tempfile.mkdtemp()) / "sales_report_synthetic.xlsx"
    wb.save(path)
    return path, "Planted (synthetic): " + sc.shocks[0].label()


def parse_period(text: str):
    start, end = text.split(":")
    return pd.Timestamp(start), pd.Timestamp(end)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("file", nargs="?")
    ap.add_argument("--metric")
    ap.add_argument("--a", help="period A as YYYY-MM-DD:YYYY-MM-DD")
    ap.add_argument("--b", help="period B as YYYY-MM-DD:YYYY-MM-DD")
    args = ap.parse_args()

    truth = None
    path = args.file
    if path is None:
        path, truth = build_messy_sample()
        print(f"Built a messy synthetic workbook: {path}\n{truth}\n")

    try:
        first = prepare(path)
    except DataError as exc:
        print(f"Cannot use this file: {exc}")
        sys.exit(1)
    if args.a and args.b:
        periods = (parse_period(args.a), parse_period(args.b))
    else:
        periods = default_periods(first)
    p = prepare(path, period_a=periods[0], period_b=periods[1]) if periods else first

    print("GUESSED COLUMN ROLES (check these before trusting anything):")
    print(p.roles.explain())
    print("\nVALIDATION REPORT:")
    print(p.report.to_text())
    if not periods:
        print("\nNo default periods (need daily data covering at least 14 days). Pass --a and --b.")
        return
    if not p.report.ok:
        print("\nStopping: fix the errors above first.")
        return

    metric = args.metric or p.metric_cols[0]
    (a0, a1), (b0, b1) = periods
    print(f"\nEXPLAINING '{metric}': {a0.date()}..{a1.date()} vs {b0.date()}..{b1.date()}")
    ex = explain_with_trust(p.data, additive(metric), (a0, a1), (b0, b1), p.dimension_cols,
                            date_col=p.date_col)
    print(f"{metric}: {ex.total_a:,.0f} -> {ex.total_b:,.0f} | VERDICT: {ex.verdict}")
    for f in ex.findings:
        print(f"  #{f.rank} {f.label():<34} explained {f.explained:+,.0f} "
              f"({f.share_of_change:.0%} of net change), "
              f"95% CI [{f.ci_low:+,.0f}, {f.ci_high:+,.0f}]")
    for note in ex.notes:
        print(f"  note: {note}")


if __name__ == "__main__":
    np.random.seed(0)
    main()
