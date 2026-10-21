"""Tests for the ingest layer. Plain asserts; run with `python -m pytest -q`."""

import io
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

from whymoved.decompose import STANDARD_METRICS
from whymoved.generator import make_scenario
from whymoved.ingest import (
    DataError,
    _parse_dates,
    _to_numeric,
    default_periods,
    guess_roles,
    prepare,
    read_table,
)
from whymoved.trust import explain_with_trust


def _raises(exc, fn, *args, **kwargs) -> bool:
    try:
        fn(*args, **kwargs)
    except exc:
        return True
    return False


def _frame(days: int = 40, regions=("N", "S"), seed: int = 0) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    start = pd.Timestamp("2026-01-05")
    rows = [
        {
            "date": start + pd.Timedelta(days=i),
            "region": r,
            "channel": c,
            "orders": int(rng.integers(5, 50)),
            "revenue": float(rng.uniform(100, 5000)),
        }
        for i in range(days)
        for r in regions
        for c in ("web", "app")
    ]
    return pd.DataFrame(rows)


def _csv(df: pd.DataFrame, **kwargs) -> Path:
    tmp = Path(tempfile.mkdtemp())
    path = tmp / kwargs.pop("name", "data.csv")
    df.to_csv(path, index=False, **kwargs)
    return path


# ------------------------------------------------------------------ reading


def test_csv_delimiters_are_detected():
    df = _frame(days=15)
    for sep in (",", ";", "\t", "|"):
        out, info = read_table(_csv(df, sep=sep))
        assert list(out.columns) == list(df.columns), sep
        assert len(out) == len(df)
        assert info["delimiter"] == sep


def test_csv_encodings_bom_and_windows():
    df = _frame(days=15)
    df["region"] = df["region"].map({"N": "Café", "S": "Zürich"})
    tmp = Path(tempfile.mkdtemp())
    (tmp / "bom.csv").write_bytes(("﻿" + df.to_csv(index=False)).encode("utf-8"))
    (tmp / "win.csv").write_bytes(df.to_csv(index=False).encode("cp1252"))
    for name in ("bom.csv", "win.csv"):
        out, _ = read_table(tmp / name)
        assert out.columns[0] == "date", (name, list(out.columns))
        assert set(out["region"]) == {"Café", "Zürich"}, name


def test_dataframe_and_uploaded_file_object_are_accepted():
    df = _frame(days=15)
    assert len(read_table(df)[0]) == len(df)
    buf = io.BytesIO(df.to_csv(index=False).encode("utf-8"))
    out, _ = read_table(buf, filename="upload.csv")
    assert len(out) == len(df)
    assert _raises(DataError, read_table, io.BytesIO(b"a,b\n1,2\n"))  # no filename


def test_unusable_files_raise_data_error():
    tmp = Path(tempfile.mkdtemp())
    assert _raises(DataError, read_table, tmp / "missing.csv")
    (tmp / "old.xls").write_bytes(b"x")
    assert _raises(DataError, read_table, tmp / "old.xls")
    (tmp / "doc.pdf").write_bytes(b"x")
    assert _raises(DataError, read_table, tmp / "doc.pdf")
    (tmp / "empty.csv").write_text("")
    assert _raises(DataError, read_table, tmp / "empty.csv")
    (tmp / "headeronly.csv").write_text("date,region,revenue\n")
    assert _raises(DataError, read_table, tmp / "headeronly.csv")
    (tmp / "ragged.csv").write_text("a,b,c\n1,2,3\n4,5,6,7,8\n")
    assert _raises(DataError, read_table, tmp / "ragged.csv")


# ------------------------------------------------------------------ coercion


def test_numbers_with_currency_indian_commas_brackets_percent():
    s = pd.Series(["₹1,23,456", "(500)", "Rs. 2,000", "1.5e3", "12%", "abc", None, " 7 "])
    vals, frac = _to_numeric(s)
    assert vals.iloc[0] == 123456 and vals.iloc[1] == -500 and vals.iloc[2] == 2000
    assert vals.iloc[3] == 1500 and abs(vals.iloc[4] - 0.12) < 1e-12 and vals.iloc[7] == 7
    assert np.isnan(vals.iloc[5]) and np.isnan(vals.iloc[6])
    assert abs(frac - 6 / 7) < 1e-12  # 'abc' failed, the blank is not counted


def test_date_order_is_inferred_or_flagged_as_ambiguous():
    day_first, i1 = _parse_dates(pd.Series(["25/12/2026", "01/02/2026"]))
    assert day_first.iloc[0] == pd.Timestamp("2026-12-25")
    assert day_first.iloc[1] == pd.Timestamp("2026-02-01") and not i1["ambiguous"]

    month_first, i2 = _parse_dates(pd.Series(["12/25/2026", "02/01/2026"]))
    assert month_first.iloc[0] == pd.Timestamp("2026-12-25")
    assert month_first.iloc[1] == pd.Timestamp("2026-02-01") and not i2["ambiguous"]

    amb, i3 = _parse_dates(pd.Series(["01/02/2026", "03/04/2026"]))
    assert i3["ambiguous"] and amb.iloc[0] == pd.Timestamp("2026-02-01")  # day first assumed

    iso, i4 = _parse_dates(pd.Series(["2026-01-05", "2026-02-03"]))
    assert iso.iloc[1] == pd.Timestamp("2026-02-03") and not i4["ambiguous"]

    refused, _ = _parse_dates(pd.Series(["20260105", "20260106"]))
    assert refused.isna().all()  # bare digits could be anything, so it never guesses


# ------------------------------------------------------------------ roles


def test_roles_on_generator_data_and_junk_columns():
    sc = make_scenario("single", 7)
    d = sc.data.head(2000).reset_index(drop=True)
    d["row_id"] = range(1, len(d) + 1)
    d["order_ref"] = np.random.default_rng(0).permutation(len(d)) + 100000
    d["notes"] = [f"free text {i}" for i in range(len(d))]
    d["source"] = "web-export"
    d["empty"] = np.nan
    roles = guess_roles(d)
    assert roles.date == "date"
    assert set(roles.dimensions) == {"region", "channel", "device", "category"}
    assert {"visits", "orders", "revenue", "cost"} <= set(roles.metrics)
    ignored = {g.name for g in roles.guesses if g.role == "ignore"}
    assert {"row_id", "order_ref", "notes", "source", "empty"} <= ignored
    assert all(g.reason for g in roles.guesses)


def test_calendar_column_becomes_a_category_and_extra_date_is_ignored():
    d = _frame(days=40)
    d["year"] = 2026
    d["month"] = d["date"].dt.month
    d["shipped"] = d["date"] + pd.Timedelta(days=2)
    roles = guess_roles(d)
    assert "month" in roles.dimensions
    assert roles.date == "date"
    assert any(g.name == "shipped" and g.role == "ignore" for g in roles.guesses)


# ------------------------------------------------------------------ cleaning and validation


def test_category_spelling_variants_merge_and_are_reported():
    d = _frame(days=30)
    d.loc[d.index % 7 == 0, "region"] = "n "
    d.loc[d.index % 11 == 0, "region"] = "  N"
    p = prepare(_csv(d))
    assert set(p.data["region"]) == {"N", "S"}
    assert p.report.has("category_variants")
    assert p.data["region"].value_counts().sum() == len(d)


def test_total_rows_removed_so_sums_do_not_double_count():
    d = _frame(days=30)
    true_total = d["revenue"].sum()
    totals = pd.DataFrame([
        {"region": "Total", "revenue": true_total, "orders": d["orders"].sum()},
        {"date": "Grand Total", "revenue": true_total},
    ])
    p = prepare(_csv(pd.concat([d, totals], ignore_index=True)))
    assert abs(p.data["revenue"].sum() - true_total) < 1e-6
    assert p.report.get("total_rows_removed")[0].count == 2


def test_duplicates_are_reported_but_kept_unless_asked():
    d = _frame(days=30)
    dup = pd.concat([d, d.head(10)], ignore_index=True)
    kept = prepare(_csv(dup))
    assert len(kept.data) == len(dup) and kept.report.get("duplicate_rows")[0].count == 10
    dropped = prepare(_csv(dup), drop_duplicates=True)
    assert len(dropped.data) == len(d) and dropped.report.has("duplicates_removed")


def test_bad_and_missing_dates_are_dropped_and_reported():
    d = _frame(days=60)
    d["date"] = d["date"].dt.strftime("%Y-%m-%d")
    d.loc[0, "date"] = "n/a"
    d.loc[1, "date"] = "TBD"
    d.loc[2, "date"] = None
    p = prepare(_csv(d))
    assert len(p.data) == len(d) - 3
    # pandas reads "n/a" as blank, so: "TBD" is unparsable; "n/a" and the empty cell are missing.
    assert p.report.get("unparsable_dates")[0].examples == ["TBD"]
    assert p.report.get("missing_dates")[0].count == 2
    assert p.report.ok  # 3 of 240 rows is under the 5% limit


def test_too_many_bad_dates_is_an_error_not_a_crash():
    d = _frame(days=60)
    d["date"] = d["date"].dt.strftime("%Y-%m-%d")
    d.loc[d.index % 4 == 0, "date"] = "unknown"
    p = prepare(_csv(d))
    assert not p.report.ok and p.report.has("too_many_bad_dates")


def test_date_gaps_are_flagged():
    d = _frame(days=40)
    d = d[~d["date"].isin([pd.Timestamp("2026-01-15"), pd.Timestamp("2026-01-16")])]
    p = prepare(_csv(d))
    gap = p.report.get("date_gaps")[0]
    assert gap.count == 2 and "2026-01-15" in gap.examples


def test_negative_non_numeric_and_blank_metric_values():
    d = _frame(days=30).astype({"revenue": object})
    d.loc[0, "revenue"] = -250.0
    d.loc[1, "revenue"] = "oops"
    d.loc[2, "revenue"] = None
    p = prepare(_csv(d))
    assert p.data.loc[0, "revenue"] == -250.0  # kept: may be a refund
    assert p.data.loc[1, "revenue"] == 0.0 and p.data.loc[2, "revenue"] == 0.0
    assert p.report.get("negative_values")[0].count == 1
    assert p.report.get("non_numeric_values")[0].examples == ["oops"]
    assert p.report.get("missing_values")[0].count == 1


def test_rare_categories_group_into_other_and_many_categories_are_capped():
    rng = np.random.default_rng(3)
    d = _frame(days=60, regions=("N", "S", "E"))
    rare = rng.random(len(d)) < 0.004
    d.loc[rare, "region"] = [f"tiny{i}" for i in range(int(rare.sum()))]
    p = prepare(_csv(d), min_category_share=0.01)
    assert "Other" in set(p.data["region"]) and p.report.has("rare_categories_grouped")

    d2 = _frame(days=60)
    d2["sku"] = [f"sku{i % 40:02d}" for i in range(len(d2))]
    p2 = prepare(_csv(d2), dimension_cols=["region", "sku"], max_categories=10)
    assert p2.data["sku"].nunique() == 10 and "Other" in set(p2.data["sku"])
    assert p2.report.has("categories_capped")


def test_merged_cell_style_blanks_are_flagged_and_fill_down_fixes_them():
    d = _frame(days=30).sort_values(["region", "date"]).reset_index(drop=True)
    shown = d.copy()
    shown.loc[shown["region"] == shown["region"].shift(), "region"] = None
    p = prepare(_csv(shown))
    blank = p.report.get("blank_categories")[0]
    assert "merged-cell" in blank.message and "Unknown" in set(p.data["region"])
    fixed = prepare(_csv(shown), fill_down=["region"])
    assert not fixed.report.has("blank_categories")
    assert set(fixed.data["region"]) == {"N", "S"}


# ------------------------------------------------------------------ excel


def _messy_workbook() -> Path:
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Read me first"
    ws = wb.create_sheet("Sales")
    ws["A1"] = "Quarterly sales report"
    ws["A2"] = "Prepared by finance"
    header = ["Date", "Region", "Channel", "Orders", "Revenue"]
    for j, h in enumerate(header, start=1):
        ws.cell(row=4, column=j, value=h)
    r = 5
    rng = np.random.default_rng(1)
    for i in range(30):
        for region in ("North", "South"):
            when = pd.Timestamp("2026-01-05") + pd.Timedelta(days=i, hours=9)
            ws.cell(row=r, column=1, value=when)
            ws.cell(row=r, column=2, value=region)
            ws.cell(row=r, column=3, value="web")
            ws.cell(row=r, column=4, value=int(rng.integers(5, 50)))
            ws.cell(row=r, column=5, value=f"₹{rng.integers(1000, 9000):,}")
            r += 1
    ws.cell(row=r, column=2, value="Total")
    ws.cell(row=r, column=4, value=999)
    path = Path(tempfile.mkdtemp()) / "report.xlsx"
    wb.save(path)
    return path


def test_excel_with_title_lines_blank_rows_total_row_and_text_money():
    p = prepare(_messy_workbook())
    assert p.report.has("header_found") and p.report.has("sheet_used")
    assert p.info["sheet"] == "Sales"
    assert p.date_col == "Date" and set(p.dimension_cols) == {"Region"}  # Channel is constant
    assert set(p.metric_cols) == {"Orders", "Revenue"}
    assert len(p.data) == 60 and p.report.has("total_rows_removed")
    assert p.data["Revenue"].min() >= 1000  # '₹4,500' text became a number
    assert (p.data["Date"] == p.data["Date"].dt.normalize()).all()  # 09:00 times dropped
    assert p.report.has("times_dropped")


def test_excel_sheet_selection_errors_are_clear():
    path = _messy_workbook()
    assert prepare(path, sheet="Sales").info["sheet"] == "Sales"
    assert _raises(DataError, prepare, path, sheet="Nope")
    assert _raises(DataError, prepare, path, sheet=9)
    assert _raises(DataError, prepare, path, sheet="Cover")  # one line only


# ------------------------------------------------------------------ fatal problems


def test_missing_roles_raise_data_error_with_a_reason():
    d = _frame(days=30)
    assert _raises(DataError, prepare, _csv(d.drop(columns=["date"])))
    assert _raises(DataError, prepare, _csv(d.drop(columns=["orders", "revenue"])))
    assert _raises(DataError, prepare, _csv(d.drop(columns=["region", "channel"])))
    assert _raises(DataError, prepare, _csv(d), date_col="nope")
    try:
        prepare(_csv(d.drop(columns=["date"])))
    except DataError as exc:
        assert "no date column" in str(exc)


def test_user_can_override_the_guessed_roles():
    d = _frame(days=30)
    p = prepare(_csv(d), metric_cols=["revenue"], dimension_cols=["region"])
    assert p.metric_cols == ["revenue"] and p.dimension_cols == ["region"]
    assert list(p.data.columns) == ["date", "region", "revenue"]


# ------------------------------------------------------------------ periods


def test_period_checks():
    sc = make_scenario("single", 7)
    d = sc.data.head(6000)
    good = prepare(_csv(sc.data), period_a=sc.period_a, period_b=sc.period_b)
    for code in ("unequal_period_lengths", "not_whole_weeks", "periods_overlap", "empty_period"):
        assert not good.report.has(code), code
    assert good.report.ok

    a0, a1 = sc.period_a
    b0, b1 = sc.period_b
    overlap = prepare(_csv(d), period_a=(a0, a1), period_b=(a1 - pd.Timedelta(days=3), b1))
    assert overlap.report.has("periods_overlap") and not overlap.report.ok

    unequal = prepare(_csv(sc.data), period_a=(a0, a1), period_b=(b0, b0 + pd.Timedelta(days=20)))
    assert unequal.report.has("unequal_period_lengths")

    ten = prepare(_csv(sc.data), period_a=(a0, a0 + pd.Timedelta(days=9)),
                  period_b=(b0, b0 + pd.Timedelta(days=9)))
    assert ten.report.has("not_whole_weeks")

    beyond = prepare(_csv(sc.data), period_a=(a0, a1), period_b=(b0, b1 + pd.Timedelta(days=30)))
    assert beyond.report.has("partial_period")

    far = prepare(_csv(sc.data), period_a=(pd.Timestamp("2030-01-01"), pd.Timestamp("2030-01-28")),
                  period_b=(b0, b1))
    assert far.report.has("empty_period") and not far.report.ok


def test_categories_that_appear_or_vanish_between_periods_are_reported():
    d = _frame(days=56)
    d.loc[(d["date"] >= pd.Timestamp("2026-02-02")) & (d["region"] == "S"), "region"] = "S2"
    p = prepare(
        _csv(d),
        period_a=(pd.Timestamp("2026-01-05"), pd.Timestamp("2026-02-01")),
        period_b=(pd.Timestamp("2026-02-02"), pd.Timestamp("2026-03-01")),
    )
    new, gone = p.report.get("new_categories")[0], p.report.get("vanished_categories")[0]
    assert new.examples == ["S2"] and gone.examples == ["S"]


def test_default_periods_are_equal_adjacent_and_whole_weeks():
    p = prepare(_csv(make_scenario("none", 3).data))
    (a0, a1), (b0, b1) = default_periods(p)
    assert (a1 - a0) == (b1 - b0) and (a1 - a0).days + 1 == 28
    assert b0 - a1 == pd.Timedelta(days=1) and b1 == p.info["date_max"]
    short = prepare(_csv(_frame(days=20)))
    assert default_periods(short)[0][1] < default_periods(short)[1][0]
    assert default_periods(prepare(_csv(_frame(days=10)))) is None


# ------------------------------------------------------------------ end to end


def test_messy_file_still_finds_the_planted_cause():
    spec = STANDARD_METRICS["revenue"]
    hits = 0
    for seed in (7, 8, 9):
        sc = make_scenario("single", seed)
        d = sc.data.copy().reset_index(drop=True)
        d["date"] = d["date"].dt.strftime("%d/%m/%Y")
        d["invoice_id"] = range(1, len(d) + 1)
        loud = d.index % 40 == 0
        d.loc[loud, "region"] = d.loc[loud, "region"].str.upper() + " "
        d["revenue"] = d["revenue"].map(lambda v: f"₹{v:,.2f}")
        d = pd.concat([d, pd.DataFrame([{"region": "Grand Total"}])], ignore_index=True)
        p = prepare(_csv(d, sep=";"), period_a=sc.period_a, period_b=sc.period_b)
        assert p.report.ok and "invoice_id" not in p.data.columns
        clean_total = p.data.loc[p.data["date"].between(*sc.period_a), "revenue"].sum()
        truth_total = sc.data.loc[sc.data["date"].between(*sc.period_a), "revenue"].sum()
        assert abs(clean_total - truth_total) < 1.0  # money parsed to the paisa, no double count
        ex = explain_with_trust(p.data, spec, sc.period_a, sc.period_b, p.dimension_cols,
                                date_col=p.date_col)
        hits += bool(ex.findings) and ex.findings[0].filters == sc.truth[0]
    assert hits == 3


def test_ready_made_rates_are_flagged_because_they_cannot_be_split():
    d = _frame(days=30)
    d["conversion_rate"] = np.random.default_rng(0).uniform(0.01, 0.09, len(d))
    d["avg_basket"] = np.random.default_rng(1).uniform(10, 90, len(d))
    p = prepare(_csv(d))
    flagged = {i.column for i in p.report.get("looks_like_rate")}
    assert flagged == {"conversion_rate", "avg_basket"}
    assert "orders / visits" in p.report.get("looks_like_rate")[0].message
    assert not prepare(_csv(_frame(days=30))).report.has("looks_like_rate")


def test_decimal_amounts_are_metrics_not_dates():
    """Regression: a text column of amounts like '1023.45' was once parsed as a date and then
    ignored as an 'extra date column'."""
    rng = np.random.default_rng(1)
    n = 40
    csv = "date,region,revenue,fee\n" + "\n".join(
        f"2026-01-{1 + i % 28:02d},{'N' if i % 2 else 'S'},{1000 + rng.uniform(0, 60):.2f},"
        f"{chr(34) + '1,234.50' + chr(34) if i % 3 else '99.5'}" for i in range(n))
    p = prepare(io.BytesIO(csv.encode()), filename="a.csv")
    assert p.date_col == "date" and "revenue" in p.metric_cols and "fee" in p.metric_cols
    parsed, _ = _parse_dates(pd.Series(["1023.45", "1031.2", "1050.00"], dtype="string"))
    assert parsed.isna().all()
    # real dotted dates still work
    parsed, _ = _parse_dates(pd.Series(["05.01.2026", "06.01.2026", "07.01.2026"], dtype="string"))
    assert parsed.notna().all()
